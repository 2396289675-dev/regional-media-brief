#!/usr/bin/env python3
"""
区域媒体情况简报 markdown -> docx 渲染器

用法:
    python build_docx.py brief.md ~/Downloads/浙江媒体情况简报.docx

支持的 markdown 语法（仅此几种）:
    # 标题1（居中大标题，全篇一次）
    ## 标题2
    ### 标题3
    **加粗**（段首标签用）
    ==红色加粗==（现场重点，控制用量）
    | a | b |  连续行 -> 表格，首行为加粗表头

版式与《陕西媒体情况简报》定版一致：黑体标题层级、宋体 11pt 正文、1.15 行距、
Table Grid 全边框表格、页脚居中页码。
"""

import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

SONG = "宋体"
HEI = "黑体"
RED = RGBColor(0xFF, 0x00, 0x00)


# --------------------------------------------------------------------------
# 基础工具
# --------------------------------------------------------------------------
def set_run_font(run, name=SONG, size=11, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:cs"), SONG)
    rFonts.set(qn("w:hint"), "eastAsia")
    return run


def emit_runs(paragraph, text, size=11, name=SONG):
    """把含 **bold** / ==red bold== 的文本切成 run 写入段落。"""
    i = 0
    n = len(text)
    while i < n:
        matched = False
        for mark, kwargs in (
            ("**", {"bold": True}),
            ("==", {"bold": True, "color": RED}),
        ):
            if text.startswith(mark, i):
                j = text.find(mark, i + 2)
                if j != -1 and j > i + 2:
                    inner = text[i + 2 : j]
                    if inner.strip():
                        set_run_font(
                            paragraph.add_run(inner), name=name, size=size, **kwargs
                        )
                    i = j + 2
                    matched = True
                    break
        if matched:
            continue
        nxt = n
        for mark in ("**", "=="):
            k = text.find(mark, i)
            if k != -1:
                nxt = min(nxt, k)
        if nxt <= i:
            nxt = i + 1
        set_run_font(paragraph.add_run(text[i:nxt]), name=name, size=size)
        i = nxt


def is_table_row(line):
    s = line.strip()
    return s.startswith("|") and s.endswith("|") and len(s) > 1


def is_separator_row(line):
    return bool(re.fullmatch(r"\|[\s:\-|]+\|", line.strip()))


def split_row(line):
    s = line.strip().strip("|")
    return [c.strip() for c in s.split("|")]


# --------------------------------------------------------------------------
# 文档骨架
# --------------------------------------------------------------------------
def setup_document(doc):
    normal = doc.styles["Normal"]
    normal.font.name = SONG
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), SONG)
    pf = normal.paragraph_format
    pf.line_spacing = 1.15
    pf.space_after = Pt(10)

    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = Mm(25.4)
    sec.bottom_margin = Mm(25.4)
    sec.left_margin = Mm(31.8)
    sec.right_margin = Mm(31.8)


def add_heading(doc, text, level):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if level == 1:
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.space_before = Pt(8.4)
        pf.space_after = Pt(13.2)
        pf.line_spacing = 1.43
        size = 22
    elif level == 2:
        pf.space_before = Pt(13.2)
        pf.space_after = Pt(10)
        pf.line_spacing = 1.15
        size = 14
    else:
        pf.space_before = Pt(8.4)
        pf.space_after = Pt(10)
        pf.line_spacing = 1.15
        size = 12
    pf.keep_with_next = True
    set_run_font(p.add_run(text), name=HEI, size=size, bold=True)
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    emit_runs(p, text)
    return p


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=len(rows), cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    tblPr = table._tbl.tblPr
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), "5000")
    tblW.set(qn("w:type"), "pct")
    tblPr.append(tblW)

    for c, text in enumerate(header):
        cell = table.cell(0, c)
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        emit_runs(p, text, size=11)
        for r in p.runs:
            r.bold = True

    for ri, row in enumerate(body, start=1):
        for c in range(len(header)):
            text = row[c] if c < len(row) else ""
            cell = table.cell(ri, c)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            emit_runs(p, text, size=11)
    return table


def add_page_footer(section):
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    set_run_font(run, name=SONG, size=9)

    f1 = OxmlElement("w:fldChar")
    f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = "PAGE"
    f2 = OxmlElement("w:fldChar")
    f2.set(qn("w:fldCharType"), "end")
    run._r.append(f1)
    run._r.append(it)
    run._r.append(f2)


# --------------------------------------------------------------------------
# markdown -> docx
# --------------------------------------------------------------------------
def render(md_text, out_path):
    doc = Document()
    setup_document(doc)

    buf = []          # 累积正文段落
    table_rows = []   # 累积表格行

    def flush_paragraph():
        if buf:
            text = "".join(buf).strip()
            if text:
                add_body(doc, text)
            buf.clear()

    def flush_table():
        if table_rows:
            add_table(doc, table_rows)
            doc.add_paragraph()  # 表格后留一个空段，避免与下一段粘连
            table_rows.clear()

    for raw in md_text.splitlines():
        line = raw.rstrip()
        stripped = line.strip()

        if not stripped:
            flush_paragraph()
            flush_table()
            continue

        if is_table_row(stripped):
            flush_paragraph()
            if is_separator_row(stripped):
                continue
            table_rows.append(split_row(stripped))
            continue

        flush_table()

        if stripped.startswith("### "):
            flush_paragraph()
            add_heading(doc, stripped[4:].strip(), 3)
        elif stripped.startswith("## "):
            flush_paragraph()
            add_heading(doc, stripped[3:].strip(), 2)
        elif stripped.startswith("# "):
            flush_paragraph()
            add_heading(doc, stripped[2:].strip(), 1)
        else:
            buf.append(stripped)

    flush_paragraph()
    flush_table()

    add_page_footer(doc.sections[0])

    out_path = os.path.expanduser(out_path)
    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    doc.save(out_path)
    return out_path


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    src, dst = sys.argv[1], sys.argv[2]
    with open(os.path.expanduser(src), encoding="utf-8") as f:
        md_text = f.read()
    path = render(md_text, dst)
    print("saved:", path)


if __name__ == "__main__":
    main()
